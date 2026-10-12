"""Excel (XLSX) parser for vocabulary imports."""

import io
from typing import Any, Dict, List, Optional, Tuple
import openpyxl


def parse_xlsx_data(
    data_bytes: bytes, sheet_name: Optional[str] = None
) -> Tuple[List[str], List[Dict[str, str]]]:
    """
    Parse XLSX bytes into header list and list of row dicts.
    """
    file_stream = io.BytesIO(data_bytes)
    wb = openpyxl.load_workbook(file_stream, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
    else:
        sheet = wb.active

    headers: List[str] = []
    rows: List[Dict[str, str]] = []

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        str_cells = [str(c).strip() if c is not None else "" for c in row]

        if not headers:
            headers = str_cells
            continue

        row_dict: Dict[str, str] = {}
        for col_idx, col_name in enumerate(headers):
            if not col_name:
                continue
            if col_idx < len(str_cells):
                row_dict[col_name] = str_cells[col_idx]
            else:
                row_dict[col_name] = ""
        rows.append(row_dict)

    return headers, rows
